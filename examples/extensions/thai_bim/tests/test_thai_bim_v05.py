import importlib,sys,types,math,unittest,tempfile
from pathlib import Path
package=types.ModuleType('tb05pure');package.__path__=[str(Path(__file__).resolve().parents[1])];sys.modules['tb05pure']=package
E=importlib.import_module('tb05pure.engine');D=importlib.import_module('tb05pure.detailing')

def volume(faces):return sum(E.dot(f[0],E.cross(f[i],f[i+1]))/6 for f in faces for i in range(1,len(f)-1))

class DetailingTests(unittest.TestCase):
    def test_analytic_length_and_tangency(self):
        b=D.bent_bar('test',[(0,0,0),(1,0,0),(1,1,0)],.012,.024,'L-90')
        r=.03;expected=2-2*r+math.pi*r/2
        self.assertAlmostEqual(b['quantity'],expected,places=10)
        self.assertEqual(b['bbs']['bend_degrees'],[90])
        self.assertGreater(volume(b['faces']),0)
        chord=sum(E.norm(E.sub(v,u)) for u,v in zip(b['bar_path'],b['bar_path'][1:]))
        self.assertLess(chord,expected);self.assertLess(expected-chord,.0001)
        shifted=D.bent_bar('other',[(4,7,2),(5,7,2),(5,8,2)],.012,.024,'L-90')
        self.assertEqual(b['bbs']['mark'],shifted['bbs']['mark'])
    def test_hook_and_tie(self):
        host=E.box_spec('Footing',0,0,0,1.2,1.2,.4)['params']
        bars=D.reinforcement('Footing',host,hook_length=.08)
        self.assertTrue(all(s['bbs']['bend_degrees']==[90,90] and volume(s['faces'])>0 for s in bars))
        for kind,w,d,h in [('Column',.25,.25,3),('Beam',3,.2,.4)]:
            host=E.box_spec(kind,0,0,0,w,d,h)['params'];bars=D.reinforcement(kind,host)
            ties=[s for s in bars if s['bbs']['shape']=='Tie-135']
            self.assertTrue(ties);self.assertTrue(all(s['bbs']['bend_degrees']==[135,90,90,90,135] for s in ties))
            self.assertTrue(all(volume(s['faces'])>0 for s in bars))
    def test_physical_lap_and_extensions(self):
        host=E.box_spec('Column',0,0,0,.3,.3,3)['params']
        a=D.reinforcement('Column',host);b=D.reinforcement('Column',host,lap_length=.5,extension_start=.2,extension_end=.3)
        main_a=[s for s in a if s['bbs']['shape']=='Straight'];main_b=[s for s in b if s['bbs']['shape']=='Straight-splice']
        self.assertEqual(len(main_b),2*len(main_a))
        self.assertAlmostEqual(sum(s['quantity'] for s in main_b)-sum(s['quantity'] for s in main_a),len(main_a)*( .5+.2+.3),places=8)
        for i in range(0,len(main_b),2):
            aa,bb=main_b[i:i+2];self.assertAlmostEqual(aa['bar_path'][-1][2]-bb['bar_path'][0][2],.5)
    def test_invalid_fit(self):
        p=E.box_spec('Footing',0,0,0,1.2,1.2,.4)['params']
        with self.assertRaises(ValueError):D.reinforcement('Footing',p,hook_length=.5)
        with self.assertRaises(ValueError):D.reinforcement('Footing',p,lap_length=.3)
        with self.assertRaises(ValueError):D.bent_bar('bad',[(0,0,0),(.02,0,0),(.02,.02,0)],.012,.024)
        with self.assertRaises(ValueError):D.bent_bar('bad',[(0,0,0),(1,0,0)],.012,float('nan'))
        p=E.box_spec('Column',0,0,0,.25,.25,3)['params']
        with self.assertRaises(ValueError):D.reinforcement('Column',p,lap_length=.5,count_x=6)
    def test_mass_and_excel(self):
        b=D.bent_bar('a',[(0,0,0),(3,0,0)],.012,.024)
        self.assertAlmostEqual(b['bbs']['mass_kg'],3*math.pi*.012**2/4*7850)
        rows=[dict(id='one',host_uid='host',bbs=b['bbs']),dict(id='two',host_uid='host',bbs=b['bbs'])]
        tables=D.tables(rows,['Example exclusion'])
        self.assertEqual(tables[0][1][1][4],2)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'เหล็ก-BBS.xlsx';E.write_xlsx(path,tables=tables)
            import openpyxl
            workbook=openpyxl.load_workbook(path);self.assertEqual(workbook['BBS']['E2'].value,2);self.assertEqual(workbook['BBS']['G2'].value,6)
            self.assertEqual(workbook['Bars'].max_row,3);self.assertEqual(workbook['Basis and issues'].cell(7,1).value,'Example exclusion')

if __name__=='__main__':unittest.main()
