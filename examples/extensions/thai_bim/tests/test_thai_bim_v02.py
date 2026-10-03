import importlib.util,json,math,random,tempfile,unittest
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('tb_engine_v02',root/'engine.py')
E=importlib.util.module_from_spec(spec);spec.loader.exec_module(E)

class V02Tests(unittest.TestCase):
    def check_plan(self,runs,stock=6,lap=.15,kerf=.003):
        plan=E.stock_cut_plan(runs,stock,lap,kerf)
        actual={r['id']:[] for r in runs}
        for b in plan['stocks']:
            self.assertGreaterEqual(b['remaining'],-1e-7)
            self.assertAlmostEqual(sum(p['length'] for p in b['pieces'])+b['kerf']+b['remaining'],stock,places=7)
            for p in b['pieces']:
                self.assertEqual(p['material'],b['material']);actual[p['run_id']].append(p)
        for r in runs:
            pieces=actual[r['id']]
            self.assertAlmostEqual(sum(p['length'] for p in pieces)-lap*(len(pieces)-1),r['length'],places=7)
            self.assertTrue(all(0<p['length']<=stock+1e-8 for p in pieces))
        for s in plan['summary']:
            self.assertAlmostEqual(s['net_m']+s['lap_m']+s['kerf_m']+s['offcut_m'],s['purchase_m'],places=7)
        return plan

    def test_cut_exact_near_stock_and_long_runs(self):
        lengths=[6,5.999,6.001,12,11.85,.025,16.2,2.4,3.5]
        p=self.check_plan([dict(id=str(i),length=l,material='Batten') for i,l in enumerate(lengths)])
        self.assertTrue(p['summary'][0]['lap_m']>0)
        p=self.check_plan([dict(id='a',length=3,material='Batten'),dict(id='b',length=3,material='Rafter')])
        self.assertEqual(len(p['stocks']),2)

    def test_random_cut_mass_balance_and_excel(self):
        rng=random.Random(4)
        runs=[dict(id=str(i),length=rng.uniform(.05,20),material=['C125','PROFAST'][i%2]) for i in range(250)]
        p=self.check_plan(runs)
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'cut.xlsx';E.write_cut_xlsx(path,p,['Thai ทดสอบ'])
            w=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(w.sheetnames,['Summary','Cuts','Stock bars','Notes'])
            self.assertEqual(w['Stock bars'].max_row,len(p['stocks'])+1)
            w.close()

    def test_cut_validation(self):
        for kwargs in [dict(lap=6),dict(kerf=-1),dict(stock=0)]:
            with self.assertRaises(ValueError): E.stock_cut_plan([],**kwargs)
        with self.assertRaises(ValueError): E.stock_cut_plan([dict(id='x',length=2,material='A')]*2)

    def test_concave_plane_clips_multiple_segments(self):
        # U footprint: a line across the upper arms must remain two pieces.
        poly=[(0,0),(4,0),(4,4),(3,4),(3,1),(1,1),(1,4),(0,4)]
        self.assertEqual(E.clip_intervals(poly,(1,0),(0,1),2),[(0.,1.),(3.,4.)])
        plane=dict(id='U',vertices=[(x,y,3+.5*y) for x,y in poly])
        specs=E.plane_roof_specs([plane])
        self.assertEqual(len({s['slot'] for s in specs}),len(specs))
        self.assertTrue(any(s['slot'].endswith('-1') for s in specs if s['kind']=='Batten'))
        self.assertTrue(all(s['quantity']>0 for s in specs))

    def test_plane_validation(self):
        for vertices in [[(0,0,3),(4,0,3),(4,4,3),(0,4,3)],
                         [(0,0,3),(4,0,3),(4,4,5),(0,4,5.1)],
                         [(0,0,3),(4,4,5),(0,4,5),(4,0,3)]]:
            with self.assertRaises(ValueError): E.plane_geometry(vertices)

    def test_family10_reference_net_lengths(self):
        old=json.loads((root/'tests/fixtures/roof-reference-totals.json').read_text(encoding='utf-8'))
        # Reference plane definitions derived from known R03 roof vertices.
        sl=math.tan(math.radians(35));ze=10.65-5.5*sl
        P={'A':(-1,-1,ze),'B':(5,-1,ze),'C':(5,.8,ze),'D':(10,.8,ze),'E':(10,12.4,ze),'F':(-1,12.4,ze),'R':(4.5,6.9,10.65),'S':(4.5,6.3,10.65),'T':(2,3.8,ze+3*sl),'U':(2,2,ze+3*sl)}
        planes=[dict(id=str(i),vertices=[P[k] for k in keys]) for i,keys in enumerate(['FER','EDSR','DCTS','CBUT','BAU','AFRSTU'],1)]
        specs=E.plane_roof_specs(planes,batten_inset=.05/math.cos(math.radians(35)),batten_end_inset=.04/math.cos(math.radians(35)))
        rafters=[s for s in specs if s['kind']=='Rafter']
        self.assertEqual(len(rafters),51)
        self.assertAlmostEqual(sum(s['quantity'] for s in rafters),old['rafter_m'],places=6)
        battens=[s for s in specs if s['kind']=='Batten']
        self.assertEqual(len(battens),116)
        self.assertAlmostEqual(sum(s['quantity'] for s in battens),old['batten_m'],places=6)

if __name__=='__main__':unittest.main()
